import io
import re
import datetime
from typing import List, Dict, Any, Optional, Set
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def parse_indonesian_amount(amt_str: str) -> float:
    """Convert amount string '1.500.000,00' or '250.000,00-' to float."""
    if not amt_str:
        return 0.0
    clean = amt_str.strip().rstrip('-').replace('.', '').replace(',', '.')
    try:
        return float(clean)
    except ValueError:
        return 0.0


def normalize_bank_date(g1: str, g2: str, g3: str) -> Dict[str, Any]:
    """
    Parse (g1, g2, g3) as DD MM YY format.
    Example:
      26 02 24 -> 26/02/2024 (Day 26, Month 02, Year 2024)
      05 06 26 -> 05/06/2026 (Day 05, Month 06, Year 2026)
    """
    day = int(g1)
    month = int(g2)
    yy = int(g3)
    
    # 2-digit year (e.g. 24 -> 2024, 25 -> 2025, 26 -> 2026)
    year = 2000 + yy if yy < 100 else yy

    # Sanity bounds check
    month = max(1, min(12, month))
    day = max(1, min(31, day))

    iso_date = f"{year:04d}-{month:02d}-{day:02d}"
    display_date = f"{day:02d}/{month:02d}/{year:04d}"

    return {
        "year": year,
        "month": month,
        "day": day,
        "iso_date": iso_date,
        "display_date": display_date
    }


def parse_recap_log(raw_text: str) -> Dict[str, Any]:
    r"""
    Parse banking log text file using PRD regex specifications:
    - Block split: r'(?:\r?\n)(?=36\s+)' or r'(?=\b36\s+.*\bPEMINDAHAN\b)'
    - Account: r'PEMINDAHAN\s+KE\s+(\d+)'
    - Date: r'(?:21|01)\s+(\d{2})(\d{2})(\d{2})' (Format DD MM YY)
    - Debit Amount: r'(\d{1,3}(?:\.\d{3})*,\d{2}-)'
    """
    # Split text into blocks
    blocks = re.split(r'(?:\r?\n)(?=36\s+)|(?=\b36\s+.*\bPEMINDAHAN\b)', raw_text)
    
    account_regex = re.compile(r'PEMINDAHAN\s+KE\s+(\d+)', re.IGNORECASE)
    date_regex = re.compile(r'(?:21|01)\s+(\d{2})(\d{2})(\d{2})')

    transactions: List[Dict[str, Any]] = []
    unique_accounts: Set[str] = set()
    years_detected: Set[int] = set()

    for block in blocks:
        block_clean = block.strip()
        if not block_clean:
            continue

        acc_match = account_regex.search(block_clean)
        date_match = date_regex.search(block_clean)

        # Debit transaction amount (ends with '-')
        debit_match = re.search(r'(\d{1,3}(?:\.\d{3})*,\d{2})-', block_clean)
        if debit_match:
            amount_val = parse_indonesian_amount(debit_match.group(1))
            raw_amt_str = debit_match.group(0)
        else:
            # Fallback: capture first formatted nominal
            amt_match = re.search(r'(\d{1,3}(?:\.\d{3})*,\d{2})', block_clean)
            if amt_match:
                amount_val = parse_indonesian_amount(amt_match.group(1))
                raw_amt_str = amt_match.group(1)
            else:
                amount_val = 0.0
                raw_amt_str = ""

        if acc_match and date_match and amount_val > 0:
            acc_no = acc_match.group(1).strip()
            date_info = normalize_bank_date(date_match.group(1), date_match.group(2), date_match.group(3))

            unique_accounts.add(acc_no)
            years_detected.add(date_info["year"])
            transactions.append({
                "account_number": acc_no,
                "year": date_info["year"],
                "month": date_info["month"],
                "day": date_info["day"],
                "iso_date": date_info["iso_date"],
                "display_date": date_info["display_date"],
                "amount": amount_val,
                "raw_amount_str": raw_amt_str
            })

    # Sort transactions chronologically by date
    transactions.sort(key=lambda x: x["iso_date"])
    sorted_accounts = sorted(list(unique_accounts))
    sorted_years = sorted(list(years_detected)) or [2024, 2025, 2026]

    return {
        "total_transactions": len(transactions),
        "unique_accounts": sorted_accounts,
        "years_detected": sorted_years,
        "transactions": transactions
    }


def generate_dual_sheet_excel(
    parsed_data: Dict[str, Any],
    account_name_mapping: Dict[str, str],
    target_years: Optional[List[int]] = None
) -> bytes:
    """
    Generate styled dual-sheet Excel file (.xlsx):
    - Sheet 1: 'Rekap Transaksi' (Daily Transaction Matrix with separate sub-rows for identical txs)
    - Sheet 2: 'Rekapan Bulanan' (Multi-Year Recap per Person with =SUM formulas)
    """
    wb = openpyxl.Workbook()
    # Default sheet
    ws1 = wb.active
    ws1.title = "Rekap Transaksi"
    ws2 = wb.create_sheet(title="Rekapan Bulanan")

    transactions = parsed_data.get("transactions", [])
    unique_accounts = parsed_data.get("unique_accounts", [])
    years = target_years or parsed_data.get("years_detected", [2024, 2025, 2026])

    # Styles
    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_bold = Font(name="Calibri", size=11, bold=True)
    font_regular = Font(name="Calibri", size=11)
    
    fill_navy = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")  # Navy Header
    fill_slate = PatternFill(start_color="334155", end_color="334155", fill_type="solid") # Slate Header
    fill_total = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid") # Light Gray Total
    fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid") # Alternate row
    
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )
    double_bottom_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='double', color='0F172A')
    )

    currency_format = '#,##0'

    # =========================================================================
    # SHEET 1: REKAP TRANSAKSI (Daily Transaction Matrix)
    # =========================================================================
    # Header Row 1: Title
    ws1.cell(row=1, column=1, value="REKAP TRANSAKSI HARIAN LEMBUR").font = Font(name="Calibri", size=14, bold=True)
    ws1.row_dimensions[1].height = 25

    # Header Row 3 & 4: Multi-level header
    # Col 1: Tanggal Transaksi
    ws1.cell(row=3, column=1, value="TANGGAL").font = font_header
    ws1.cell(row=3, column=1).fill = fill_navy
    ws1.cell(row=3, column=1).alignment = Alignment(horizontal="center", vertical="center")
    ws1.cell(row=4, column=1, value="DD/MM/YYYY").font = font_header
    ws1.cell(row=4, column=1).fill = fill_slate
    ws1.cell(row=4, column=1).alignment = Alignment(horizontal="center", vertical="center")
    ws1.cell(row=3, column=1).border = thin_border
    ws1.cell(row=4, column=1).border = thin_border

    # Account Columns
    col_idx = 2
    acc_to_col: Dict[str, int] = {}
    for acc in unique_accounts:
        emp_name = account_name_mapping.get(acc, f"Karyawan ({acc})")
        acc_to_col[acc] = col_idx

        # Row 3: Employee Name
        cell_name = ws1.cell(row=3, column=col_idx, value=emp_name.upper())
        cell_name.font = font_header
        cell_name.fill = fill_navy
        cell_name.alignment = Alignment(horizontal="center", vertical="center")
        cell_name.border = thin_border

        # Row 4: Direct Account Number (No 'No. Rek:' text)
        cell_acc = ws1.cell(row=4, column=col_idx, value=str(acc))
        cell_acc.font = font_header
        cell_acc.fill = fill_slate
        cell_acc.alignment = Alignment(horizontal="center", vertical="center")
        cell_acc.border = thin_border

        col_idx += 1

    # Group transactions by iso_date -> account -> list of amounts
    # Constraint: If identical transactions occur on same date & account, do NOT sum up; display on separate sub-rows
    date_map: Dict[str, Dict[str, List[float]]] = {}
    display_date_map: Dict[str, str] = {}

    for tx in transactions:
        d = tx["iso_date"]
        display_date_map[d] = tx["display_date"]
        acc = tx["account_number"]
        amt = tx["amount"]

        if d not in date_map:
            date_map[d] = {a: [] for a in unique_accounts}
        date_map[d][acc].append(amt)

    current_row = 5
    is_alternate = False

    for d in sorted(date_map.keys()):
        acc_txs = date_map[d]
        # Find max transactions on this date across all accounts
        max_subrows = max([len(tx_list) for tx_list in acc_txs.values()] or [1])
        if max_subrows == 0:
            max_subrows = 1

        for sub_i in range(max_subrows):
            # Date cell
            cell_date = ws1.cell(
                row=current_row,
                column=1,
                value=display_date_map[d] if sub_i == 0 else f"{display_date_map[d]} (sub-{sub_i+1})"
            )
            cell_date.font = font_regular
            cell_date.alignment = Alignment(horizontal="center", vertical="center")
            cell_date.border = thin_border
            if is_alternate:
                cell_date.fill = fill_zebra

            # Each account column
            for acc in unique_accounts:
                c_idx = acc_to_col[acc]
                tx_list = acc_txs[acc]
                val = tx_list[sub_i] if sub_i < len(tx_list) else None

                cell_val = ws1.cell(row=current_row, column=c_idx, value=val)
                cell_val.font = font_regular
                cell_val.alignment = Alignment(horizontal="right", vertical="center")
                cell_val.number_format = currency_format
                cell_val.border = thin_border
                if is_alternate:
                    cell_val.fill = fill_zebra

            current_row += 1
        is_alternate = not is_alternate

    # Total Row for Sheet 1
    total_row = current_row
    cell_tot_lbl = ws1.cell(row=total_row, column=1, value="TOTAL")
    cell_tot_lbl.font = font_bold
    cell_tot_lbl.fill = fill_total
    cell_tot_lbl.alignment = Alignment(horizontal="center", vertical="center")
    cell_tot_lbl.border = double_bottom_border

    for acc in unique_accounts:
        c_idx = acc_to_col[acc]
        col_letter = get_column_letter(c_idx)
        cell_tot = ws1.cell(row=total_row, column=c_idx, value=f"=SUM({col_letter}5:{col_letter}{total_row-1})")
        cell_tot.font = font_bold
        cell_tot.fill = fill_total
        cell_tot.alignment = Alignment(horizontal="right", vertical="center")
        cell_tot.number_format = currency_format
        cell_tot.border = double_bottom_border

    # Adjust Sheet 1 Column Widths
    ws1.column_dimensions["A"].width = 22
    for acc in unique_accounts:
        c_idx = acc_to_col[acc]
        ws1.column_dimensions[get_column_letter(c_idx)].width = 24

    # =========================================================================
    # SHEET 2: REKAPAN BULANAN (Multi-Year Recap: 2024, 2025, 2026 per person)
    # =========================================================================
    month_names = ["Jan", "Feb", "Mar", "Apr", "Mei", "Jun", "Jul", "Ags", "Sep", "Okt", "Nov", "Des"]
    
    # Pre-calculate monthly totals: year -> month -> acc -> sum
    year_month_data: Dict[int, Dict[int, Dict[str, float]]] = {}
    for y in years:
        year_month_data[y] = {m: {a: 0.0 for a in unique_accounts} for m in range(1, 13)}

    for tx in transactions:
        y = tx["year"]
        m = tx["month"]
        acc = tx["account_number"]
        if y in year_month_data and m in year_month_data[y] and acc in year_month_data[y][m]:
            year_month_data[y][m][acc] += tx["amount"]

    ws2_row = 1
    for y in years:
        # Title for Year
        ws2.cell(row=ws2_row, column=1, value=f"REKAPITULASI BIAYA LEMBUR TAHUN {y}").font = Font(name="Calibri", size=13, bold=True)
        ws2_row += 1

        # Table Header
        headers = ["No", "Nama Karyawan", "No. Rekening"] + month_names + ["TOTAL TAHUNAN"]
        for c_i, h_text in enumerate(headers, 1):
            c_cell = ws2.cell(row=ws2_row, column=c_i, value=h_text)
            c_cell.font = font_header
            c_cell.fill = fill_navy
            c_cell.alignment = Alignment(horizontal="center", vertical="center")
            c_cell.border = thin_border
        ws2.row_dimensions[ws2_row].height = 24
        
        table_start_row = ws2_row + 1
        ws2_row += 1

        # Rows per Person
        for emp_idx, acc in enumerate(unique_accounts, 1):
            emp_name = account_name_mapping.get(acc, f"Karyawan {emp_idx}")
            
            # No
            c_no = ws2.cell(row=ws2_row, column=1, value=emp_idx)
            c_no.alignment = Alignment(horizontal="center", vertical="center")
            c_no.font = font_regular
            c_no.border = thin_border

            # Name
            c_name = ws2.cell(row=ws2_row, column=2, value=emp_name)
            c_name.alignment = Alignment(horizontal="left", vertical="center")
            c_name.font = font_regular
            c_name.border = thin_border

            # Account No (Pure number)
            c_acc = ws2.cell(row=ws2_row, column=3, value=str(acc))
            c_acc.alignment = Alignment(horizontal="center", vertical="center")
            c_acc.font = font_regular
            c_acc.border = thin_border

            # Months 1 to 12
            for m in range(1, 13):
                col_num = 3 + m
                m_val = year_month_data[y][m][acc]
                c_m = ws2.cell(row=ws2_row, column=col_num, value=m_val if m_val > 0 else 0)
                c_m.alignment = Alignment(horizontal="right", vertical="center")
                c_m.number_format = currency_format
                c_m.font = font_regular
                c_m.border = thin_border

            # Row Horizontal Total Formula =SUM(D_row:O_row)
            start_m_col = get_column_letter(4)
            end_m_col = get_column_letter(15)
            tot_col_letter = get_column_letter(16)
            c_row_tot = ws2.cell(row=ws2_row, column=16, value=f"=SUM({start_m_col}{ws2_row}:{end_m_col}{ws2_row})")
            c_row_tot.alignment = Alignment(horizontal="right", vertical="center")
            c_row_tot.number_format = currency_format
            c_row_tot.font = font_bold
            c_row_tot.border = thin_border

            ws2_row += 1

        table_end_row = ws2_row - 1

        # Summary Bottom Row for this Year
        ws2.cell(row=ws2_row, column=1, value="").border = double_bottom_border
        c_tot_lbl = ws2.cell(row=ws2_row, column=2, value="TOTAL KESELURUHAN")
        c_tot_lbl.font = font_bold
        c_tot_lbl.fill = fill_total
        c_tot_lbl.alignment = Alignment(horizontal="center", vertical="center")
        c_tot_lbl.border = double_bottom_border

        ws2.cell(row=ws2_row, column=3, value="").border = double_bottom_border
        ws2.cell(row=ws2_row, column=3).fill = fill_total

        # Month Columns Vertical Sum Formulas
        for col_i in range(4, 17):
            col_letter = get_column_letter(col_i)
            c_sum = ws2.cell(row=ws2_row, column=col_i, value=f"=SUM({col_letter}{table_start_row}:{col_letter}{table_end_row})")
            c_sum.font = font_bold
            c_sum.fill = fill_total
            c_sum.alignment = Alignment(horizontal="right", vertical="center")
            c_sum.number_format = currency_format
            c_sum.border = double_bottom_border

        ws2_row += 3  # Gap between year tables

    # Adjust Sheet 2 Column Widths
    ws2.column_dimensions["A"].width = 6
    ws2.column_dimensions["B"].width = 26
    ws2.column_dimensions["C"].width = 18
    for m_i in range(4, 16):
        ws2.column_dimensions[get_column_letter(m_i)].width = 15
    ws2.column_dimensions["P"].width = 20

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()
